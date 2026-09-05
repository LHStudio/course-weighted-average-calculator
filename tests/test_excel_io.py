import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_io import EXPORT_HEADERS, export_summary, read_xlsx_records


class ExcelIOTests(unittest.TestCase):
    def test_export_has_requested_layout_and_blank_uncomputable_values(self):
        rows = [
            {
                "grade_year": "2025-2026\u5927\u4e8c",
                "student_id": "001001",
                "student_name": "\u5b66\u751f\u7532",
                "total_credits": 3.5,
                "total_score": 180.0,
                "average_score": 90.0,
                "weighted_average": 91.2,
                "credit_gpa": 13.4,
                "average_gpa": 3.8286,
                "rank": 1,
            },
            {
                "grade_year": "2025-2026\u5927\u4e8c",
                "student_id": "001002",
                "student_name": "\u5b66\u751f\u4e59",
                "total_credits": None,
                "total_score": None,
                "average_score": None,
                "weighted_average": None,
                "credit_gpa": None,
                "average_gpa": None,
                "rank": None,
            },
        ]

        with tempfile.TemporaryDirectory() as directory:
            path = export_summary(Path(directory) / "result.xlsx", rows)
            workbook = load_workbook(path)
            sheet = workbook.active

            self.assertEqual([cell.value for cell in sheet[1]], list(EXPORT_HEADERS))
            self.assertEqual(sheet["B2"].value, "001001")
            self.assertEqual(sheet["D2"].value, 3.5)
            self.assertIsNone(sheet["D3"].value)
            self.assertIsNone(sheet["J3"].value)
            self.assertEqual(sheet.freeze_panes, "A2")
            self.assertEqual(sheet.auto_filter.ref, "A1:J3")
            self.assertTrue(sheet["A1"].font.bold)
            self.assertEqual(sheet["D2"].number_format, "General")
            self.assertEqual(sheet["F2"].number_format, "0.0")
            self.assertEqual(sheet["H2"].number_format, "0.0")
            self.assertEqual(sheet["I2"].number_format, "0.0000")

    def test_read_xlsx_records_uses_the_shared_record_model(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(
                [
                    "\u53d6\u5f97\u5b66\u5e74\u5b66\u671f",
                    "\u5e74\u7ea7",
                    "\u73ed\u7ea7",
                    "\u5b66\u53f7",
                    "\u59d3\u540d",
                    "\u8bfe\u7a0b\u4ee3\u7801",
                    "\u8bfe\u7a0b\u540d\u79f0",
                    "\u5b66\u5206",
                    "\u8bfe\u7a0b\u7c7b\u522b",
                    "\u53d6\u5f97\u65b9\u5f0f",
                    "\u6210\u7ee9",
                ]
            )
            sheet.append(
                [
                    "2025-2026\u5b66\u5e74\u7b2c\u4e00\u5b66\u671f",
                    "2024",
                    "\u8f6f\u4ef62402",
                    "001001",
                    "\u5b66\u751f\u7532",
                    "C-01",
                    "\u8bfe\u7a0b\u7532",
                    2,
                    "\u4e13\u4e1a\u8bfe/\u5fc5\u4fee\u8bfe",
                    "\u521d\u4fee\u53d6\u5f97",
                    "\u4f18\u79c0",
                ]
            )
            workbook.save(path)

            records = read_xlsx_records(path)

        self.assertEqual(len(records), 1)
        self.assertEqual(records[0].student_id, "001001")
        self.assertEqual(records[0].course_code, "C-01")
        self.assertEqual(records[0].raw_score, "\u4f18\u79c0")


if __name__ == "__main__":
    unittest.main()

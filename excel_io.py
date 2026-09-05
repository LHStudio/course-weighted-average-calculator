from __future__ import annotations

from pathlib import Path
from typing import Iterable, Mapping

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.dimensions import ColumnDimension

from course_calculator import CourseRecord, records_from_rows


EXPORT_HEADERS = (
    "年级",
    "学号",
    "姓名",
    "总学分",
    "总成绩",
    "平均成绩",
    "加权平均分",
    "学分绩点",
    "平均学分绩点",
    "排名",
)

EXPORT_KEYS = (
    "grade_year",
    "student_id",
    "student_name",
    "total_credits",
    "total_score",
    "average_score",
    "weighted_average",
    "credit_gpa",
    "average_gpa",
    "rank",
)


def read_xlsx_records(path: str | Path) -> list[CourseRecord]:
    workbook = load_workbook(Path(path), read_only=True, data_only=True)
    try:
        for sheet in workbook.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            try:
                return records_from_rows(rows)
            except ValueError:
                continue
    finally:
        workbook.close()
    raise ValueError("工作簿中没有找到可识别的成绩明细表。")


def export_summary(
    path: str | Path,
    rows: Iterable[Mapping[str, object]],
    title: str = "加权平均分排名",
) -> Path:
    destination = Path(path)
    if destination.suffix.lower() != ".xlsx":
        destination = destination.with_suffix(".xlsx")
    destination.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "排名结果"
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"

    sheet.append(EXPORT_HEADERS)
    written_rows = []
    for row in rows:
        values = [row.get(key) for key in EXPORT_KEYS]
        values[1] = str(values[1]) if values[1] is not None else None
        sheet.append(values)
        written_rows.append(values)

    thin_black = Side(style="thin", color="000000")
    border = Border(left=thin_black, right=thin_black, top=thin_black, bottom=thin_black)
    header_fill = PatternFill("solid", fgColor="F3F5F7")
    body_font = Font(name="Microsoft YaHei", size=10.5, color="111827")
    header_font = Font(name="Microsoft YaHei", size=11, bold=True, color="111827")

    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    sheet.row_dimensions[1].height = 34

    for row_index in range(2, sheet.max_row + 1):
        sheet.row_dimensions[row_index].height = 23
        for cell in sheet[row_index]:
            cell.font = body_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

    for column in ("D", "E"):
        for cell in sheet[column][1:]:
            cell.number_format = "General"
    for column in ("F", "G", "H"):
        for cell in sheet[column][1:]:
            cell.number_format = "0.0"
    for cell in sheet["I"][1:]:
        cell.number_format = "0.0000"
    for cell in sheet["J"][1:]:
        cell.number_format = "0"
    for cell in sheet["B"][1:]:
        cell.number_format = "@"

    widths = (18, 16, 12, 10, 12, 12, 14, 12, 16, 9)
    for column_index, width in enumerate(widths, start=1):
        letter = get_column_letter(column_index)
        sheet.column_dimensions[letter] = ColumnDimension(sheet, index=letter, width=width)

    sheet.auto_filter.ref = f"A1:J{max(sheet.max_row, 1)}"
    sheet.print_title_rows = "1:1"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddHeader.center.text = title
    sheet.oddHeader.center.size = 12
    sheet.oddHeader.center.font = "Microsoft YaHei,Bold"

    workbook.save(destination)
    workbook.close()
    return destination
